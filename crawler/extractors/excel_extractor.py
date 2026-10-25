"""
Production-grade Excel extractor for multi-sheet and flat hospital pricelists.
Encapsulates workbook parsing, column introspection, package deal extraction,
and IPD accommodation tables into validated Pydantic models.
"""

from pathlib import Path
from typing import List, Dict, Tuple, Optional, Any
import openpyxl
from crawler.extractors.base import BaseExtractor
from crawler.models import ProcedureItem
from crawler.normalizers import (
    clean_text,
    parse_numeric_price,
    format_price_string,
    sanitize_procedure_names,
    split_mixed_bilingual,
    ARABIC_CHAR_PATTERN,
)
from crawler.utils.logger import logger

SKIP_SHEETS = {"القواعد العامة", "Packages"}


class ExcelPricelistExtractor(BaseExtractor):
    """Extracts medical procedures and pricelists from legacy and modern Excel files."""

    def __init__(self, file_path: str, year: str):
        self.file_path = Path(file_path)
        self.year = year
        self.source_label = f"Elite Hospital {year}"

    def extract(self, source: Any = None) -> List[ProcedureItem]:
        """Entrypoint for extraction: routes to appropriate multi-sheet or flat parser."""
        if not self.file_path.exists():
            raise FileNotFoundError(f"Workbook not found at {self.file_path}")

        logger.info(f"Opening workbook: [cyan]{self.file_path.name}[/cyan] ({self.year})")
        wb = openpyxl.load_workbook(str(self.file_path), data_only=True)
        sheet_names = wb.sheetnames

        # Multi-sheet files (like 2023, 2026) have many specialized tabs
        if len(sheet_names) > 3:
            return self._extract_multi_sheet(wb)
        else:
            return self._extract_flat(wb)

    def _extract_multi_sheet(self, wb: openpyxl.Workbook) -> List[ProcedureItem]:
        """Parses multi-sheet departmental workbooks."""
        items: List[ProcedureItem] = []

        for ws in wb:
            title = ws.title.strip()
            if title in SKIP_SHEETS:
                continue

            en_col = ar_col = pr_col = code_col = None
            sheet_count = 0

            for row in ws.iter_rows(values_only=True):
                cells = [clean_text(v) for v in row]
                low_cells = [c.lower() for c in cells]

                # Identify header row with price column
                pcols = [i for i, c in enumerate(low_cells) if c and "price" in c and len(c) < 25 and "package" not in c]
                if pcols and not any(parse_numeric_price(v) is not None for v in row):
                    pr_col = pcols[0]
                    n_en = n_ar = n_code = None
                    for i, c in enumerate(low_cells):
                        if i == pr_col:
                            continue
                        if "arabic" in c:
                            n_ar = i
                        elif c == "code":
                            n_code = i
                        elif ("english" in c or "service name" in c) and n_en is None:
                            n_en = i

                    if n_en is None:
                        n_en = en_col if en_col is not None else next((i for i, c in enumerate(cells) if c and i != pr_col), None)
                        n_ar = n_ar if n_ar is not None else ar_col
                        n_code = n_code if n_code is not None else code_col

                    en_col, ar_col, code_col = n_en, n_ar, n_code
                    continue

                if pr_col is None or en_col is None:
                    continue

                raw_price = parse_numeric_price(row[pr_col]) if pr_col < len(row) else None
                if raw_price is None or raw_price <= 0:
                    continue

                raw_en = cells[en_col] if en_col < len(cells) else ""
                raw_ar = cells[ar_col] if ar_col is not None and ar_col < len(cells) else ""
                raw_code = cells[code_col] if code_col is not None and code_col < len(cells) else None

                en, ar, outsourced = sanitize_procedure_names(raw_en, raw_ar)
                if not (en or ar):
                    continue

                try:
                    item = ProcedureItem(
                        nameEnglish=en,
                        nameArabic=ar,
                        price=format_price_string(raw_price),
                        sectionName=title,
                        sourceFile=self.source_label,
                        code=raw_code if raw_code and raw_code not in ("#N/A", "None") else None,
                        outsourced=True if outsourced else None,
                        sources=[{"file": self.file_path.name, "sheet": title, "year": self.year}],
                        extractor="excel",
                        extractorVersion="1.2.0",
                        confidence=1.0,
                        flags=[],
                    )
                    items.append(item)
                    sheet_count += 1
                except Exception as err:
                    logger.debug(f"Row validation skipped in sheet {title}: {err}")

        # Specialized sections (IPD accommodations & cardiology/stroke packages)
        if "IPD Services" in wb.sheetnames:
            items.extend(self._extract_ipd_accommodation(wb["IPD Services"]))

        package_targets = [s for s in ["Stroke unit", "Cathlab unit", "Cathlab unit ", "Open heart"] if s in wb.sheetnames]
        if package_targets:
            items.extend(self._extract_packages(wb, package_targets))

        return items

    def _extract_ipd_accommodation(self, ws: Any) -> List[ProcedureItem]:
        """Extracts inpatient bed & room rates."""
        items: List[ProcedureItem] = []
        for row in ws.iter_rows(min_row=5, max_row=24, values_only=True):
            label = clean_text(row[0])
            nums = [parse_numeric_price(v) for v in row[1:6] if parse_numeric_price(v) is not None]
            if not label or not nums:
                continue
            total = nums[3] if len(nums) >= 4 else nums[-1]
            parts = [x.strip() for x in label.split("/")]
            ar = next((x for x in parts if ARABIC_CHAR_PATTERN.search(x)), "")
            en = next((x for x in parts if not ARABIC_CHAR_PATTERN.search(x)), "")

            try:
                items.append(
                    ProcedureItem(
                        nameEnglish=en,
                        nameArabic=ar,
                        price=format_price_string(total),
                        sectionName="IPD Services",
                        sourceFile=self.source_label,
                        sources=[{"file": self.file_path.name, "sheet": "IPD Services", "year": self.year}],
                        extractor="excel",
                        extractorVersion="1.2.0",
                        confidence=1.0,
                        flags=[],
                    )
                )
            except Exception:
                pass
        return items

    def _extract_packages(self, wb: openpyxl.Workbook, sheet_names: List[str]) -> List[ProcedureItem]:
        """Extracts package block procedures (e.g. Cathlab, Open Heart)."""
        import re
        items: List[ProcedureItem] = []
        for sname in sheet_names:
            ws = wb[sname]
            title = None
            for row in ws.iter_rows(values_only=True):
                cells = [clean_text(v) for v in row if v is not None and clean_text(v)]
                if not cells:
                    continue
                match = re.fullmatch(r"(\d+)\s*EGP", cells[0], re.IGNORECASE)
                if match:
                    name = cells[1] if len(cells) > 1 else title
                    if name:
                        en, ar = split_mixed_bilingual(name)
                        try:
                            items.append(
                                ProcedureItem(
                                    nameEnglish=en,
                                    nameArabic=ar,
                                    price=match.group(1),
                                    sectionName=sname.strip(),
                                    sourceFile=self.source_label,
                                    sources=[{"file": self.file_path.name, "sheet": sname.strip(), "year": self.year}],
                                    extractor="excel",
                                    extractorVersion="1.2.0",
                                    confidence=1.0,
                                    flags=[],
                                )
                            )
                        except Exception:
                            pass
                elif len(cells) == 1 and ARABIC_CHAR_PATTERN.search(cells[0]) and not cells[0].startswith("Package"):
                    title = cells[0]
        return items

    def _extract_flat(self, wb: openpyxl.Workbook) -> List[ProcedureItem]:
        """Extracts flat single-sheet workbooks (e.g. 2024, 2025)."""
        ws = wb.worksheets[0]
        items: List[ProcedureItem] = []

        for raw in ws.iter_rows(values_only=True):
            cells = [clean_text(c) for c in raw if c is not None and clean_text(c) != ""]
            if not cells:
                continue
            if cells[0].lower() == "english_name":
                continue

            price = None
            for c in reversed(cells):
                n = parse_numeric_price(c)
                if n is not None:
                    price = n
                    break

            if price is None or price <= 0:
                continue

            txt_cells = [c for c in cells if parse_numeric_price(c) is None]
            en_parts = [t for t in txt_cells if not ARABIC_CHAR_PATTERN.search(t)]
            ar_parts = [t for t in txt_cells if ARABIC_CHAR_PATTERN.search(t)]

            raw_en = ", ".join(en_parts) if en_parts else ""
            raw_ar = " - ".join(ar_parts) if ar_parts else ""

            en, ar, outsourced = sanitize_procedure_names(raw_en, raw_ar)
            if not (en or ar):
                continue

            try:
                items.append(
                    ProcedureItem(
                        nameEnglish=en,
                        nameArabic=ar,
                        price=format_price_string(price),
                        sectionName="General",
                        sourceFile=self.source_label,
                        outsourced=True if outsourced else None,
                        sources=[{"file": self.file_path.name, "sheet": "Sheet1", "year": self.year}],
                        extractor="excel",
                        extractorVersion="1.2.0",
                        confidence=1.0,
                        flags=[],
                    )
                )
            except Exception:
                pass

        return items
